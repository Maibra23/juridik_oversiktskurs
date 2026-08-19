"""Tester för exportmodulen: Markdownrapport och Excelarbetsbok.

Testar de rena rapportbyggarna utan Streamlit: given quizresultat och
genomförda case ska en läsbar Markdownrapport och en giltig Excelfil
(openpyxl) byggas. Session-state-läsningen testas via röktesterna.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl

from utils.export import bygg_excel_rapport, bygg_markdown_rapport

QUIZRESULTAT = {
    "Avtalsrätt": {"avt-mc-1": True, "avt-mc-2": False, "avt-mc-3": True},
    "Skadeståndsrätt": {"ska-mc-1": True},
}
CASE_GENOMFORDA = {
    "Avtalsrätt": ("avt-case-1",),
    "Straffrätt och processrätt": ("str-case-1",),
}


# --- Markdown -----------------------------------------------------------------

def test_markdown_innehaller_moduler_och_poang():
    rapport = bygg_markdown_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    assert "Avtalsrätt" in rapport
    assert "2/3" in rapport
    assert "Skadeståndsrätt" in rapport
    assert "1/1" in rapport


def test_markdown_innehaller_genomforda_case():
    rapport = bygg_markdown_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    assert "avt-case-1" in rapport
    assert "Straffrätt och processrätt" in rapport


def test_markdown_tom_data_ger_vanligt_meddelande():
    rapport = bygg_markdown_rapport({}, {})
    assert "Inga" in rapport or "inga" in rapport


def test_markdown_har_rubrik_och_disclaimer():
    rapport = bygg_markdown_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    assert rapport.startswith("# ")
    assert "juridisk rådgivning" in rapport


# --- Excel ---------------------------------------------------------------------

def test_excel_ar_giltig_arbetsbok_med_tva_blad():
    data = bygg_excel_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    wb = openpyxl.load_workbook(BytesIO(data))
    assert "Quizresultat" in wb.sheetnames
    assert "Genomförda case" in wb.sheetnames


def test_excel_quizblad_har_moduler_och_poang():
    data = bygg_excel_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    ws = openpyxl.load_workbook(BytesIO(data))["Quizresultat"]
    rader = [tuple(c.value for c in rad) for rad in ws.iter_rows()]
    assert rader[0] == ("Modul", "Rätt", "Besvarade", "Andel")
    modulrader = {r[0]: r for r in rader[1:]}
    assert modulrader["Avtalsrätt"][1] == 2
    assert modulrader["Avtalsrätt"][2] == 3
    assert modulrader["Skadeståndsrätt"][1] == 1


def test_excel_caseblad_listar_case():
    data = bygg_excel_rapport(QUIZRESULTAT, CASE_GENOMFORDA)
    ws = openpyxl.load_workbook(BytesIO(data))["Genomförda case"]
    celler = [c.value for rad in ws.iter_rows() for c in rad]
    assert "avt-case-1" in celler
    assert "str-case-1" in celler


def test_excel_tom_data_ger_giltig_arbetsbok():
    data = bygg_excel_rapport({}, {})
    wb = openpyxl.load_workbook(BytesIO(data))
    assert "Quizresultat" in wb.sheetnames
